from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from ..config import settings
from ..database import db
from . import federated_data

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".csv",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
}


class KnowledgeImportError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def text(self) -> str:
        return "\n".join(self.parts)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise KnowledgeImportError(f"Unsupported file type. Supported types: {allowed}")

    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix == ".docx":
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    elif suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(_decode_text(data))
        text = parser.text()
    elif suffix == ".json":
        raw = _decode_text(data)
        try:
            parsed = json.loads(raw)
            text = json.dumps(parsed, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            text = raw
    else:
        text = _decode_text(data)

    normalized = _normalize_text(text)
    if not normalized:
        raise KnowledgeImportError("No readable text could be extracted from this file.")
    return normalized


def _chunk_text(text: str, chunk_size: int = 1800, overlap: int = 200) -> list[tuple[int, int, int, str]]:
    text = text.strip()
    if not text:
        return [(0, 0, 0, "")]

    chunks: list[tuple[int, int, int, str]] = []
    start = 0
    index = 0
    length = len(text)

    while start < length:
        ideal_end = min(start + chunk_size, length)
        end = ideal_end

        if ideal_end < length:
            search_start = min(ideal_end, start + int(chunk_size * 0.55))
            window = text[search_start:ideal_end]
            candidates = [
                window.rfind("\n\n"),
                window.rfind("\n"),
                window.rfind(". "),
                window.rfind(" "),
            ]
            cut = max(candidates)
            if cut > 0:
                end = search_start + cut + 1

        if end <= start:
            end = min(start + chunk_size, length)

        chunk = text[start:end].strip()
        if chunk:
            chunks.append((index, start, end, chunk))
            index += 1

        if end >= length:
            break
        start = max(end - overlap, start + 1)

    return chunks or [(0, 0, 0, "")]


def _replace_chunks(connection, item_id: int, title: str, content: str) -> int:
    connection.execute("DELETE FROM knowledge_chunks WHERE knowledge_item_id=?", (item_id,))
    chunks = _chunk_text(content)
    for chunk_index, char_start, char_end, chunk in chunks:
        connection.execute(
            """
            INSERT INTO knowledge_chunks(knowledge_item_id, chunk_index, char_start, char_end, content)
            VALUES (?, ?, ?, ?, ?)
            """,
            (item_id, chunk_index, char_start, char_end, chunk),
        )
    return len(chunks)


def ensure_knowledge_index() -> int:
    indexed = 0
    with db() as connection:
        rows = connection.execute(
            """
            SELECT ki.id, ki.title, COALESCE(ki.content, '') AS content
            FROM knowledge_items ki
            WHERE NOT EXISTS (
                SELECT 1 FROM knowledge_chunks kc WHERE kc.knowledge_item_id=ki.id
            )
            ORDER BY ki.id
            """
        ).fetchall()
        for row in rows:
            _replace_chunks(connection, row["id"], row["title"], row["content"])
            indexed += 1
    return indexed


def rebuild_knowledge_index() -> dict[str, int]:
    with db() as connection:
        rows = connection.execute(
            "SELECT id, title, COALESCE(content, '') AS content FROM knowledge_items ORDER BY id"
        ).fetchall()
        total_chunks = 0
        for row in rows:
            total_chunks += _replace_chunks(connection, row["id"], row["title"], row["content"])
    return {"items": len(rows), "chunks": total_chunks}


def create_knowledge_item(
    title: str,
    kind: str,
    content: str,
    source_path: str | None = None,
) -> dict[str, Any]:
    normalized = _normalize_text(content)
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest() if normalized else None
    with db() as connection:
        cursor = connection.execute(
            """
            INSERT INTO knowledge_items(title, kind, source_path, content, content_hash)
            VALUES (?, ?, ?, ?, ?)
            """,
            (title.strip(), kind.strip().lower(), source_path, normalized, digest),
        )
        item_id = int(cursor.lastrowid)
        chunk_count = _replace_chunks(connection, item_id, title.strip(), normalized)
    return {"created": True, "id": item_id, "chunk_count": chunk_count}


def ingest_document(filename: str, media_type: str | None, data: bytes) -> dict[str, Any]:
    original_name = Path(filename or "").name.strip()
    if not original_name:
        raise KnowledgeImportError("A file name is required.")
    if len(data) > settings.max_upload_bytes:
        raise KnowledgeImportError(
            f"File exceeds the {settings.max_upload_bytes // (1024 * 1024)} MB upload limit.",
            status_code=413,
        )

    suffix = Path(original_name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise KnowledgeImportError(f"Unsupported file type. Supported types: {allowed}")

    file_hash = hashlib.sha256(data).hexdigest()
    with db() as connection:
        duplicate = connection.execute(
            """
            SELECT ki.id, ki.title, kd.original_name
            FROM knowledge_items ki
            JOIN knowledge_documents kd ON kd.knowledge_item_id=ki.id
            WHERE ki.kind='document' AND ki.content_hash=?
            LIMIT 1
            """,
            (file_hash,),
        ).fetchone()
    if duplicate is not None:
        return {
            "created": False,
            "duplicate": True,
            "id": duplicate["id"],
            "title": duplicate["title"],
            "original_name": duplicate["original_name"],
        }

    try:
        extracted = extract_text(original_name, data)
    except KnowledgeImportError:
        raise
    except Exception as exc:
        raise KnowledgeImportError(f"Unable to read {original_name}: {exc}") from exc

    settings.knowledge_files_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}{suffix}"
    target = settings.knowledge_files_dir / stored_name
    target.write_bytes(data)

    title = Path(original_name).stem.strip()[:240] or original_name[:240]
    metadata = {
        "original_name": original_name,
        "media_type": media_type or "",
        "size_bytes": len(data),
    }

    try:
        with db() as connection:
            cursor = connection.execute(
                """
                INSERT INTO knowledge_items(title, kind, source_path, content, content_hash, metadata_json)
                VALUES (?, 'document', ?, ?, ?, ?)
                """,
                (
                    title,
                    original_name,
                    extracted,
                    file_hash,
                    json.dumps(metadata, separators=(",", ":")),
                ),
            )
            item_id = int(cursor.lastrowid)
            connection.execute(
                """
                INSERT INTO knowledge_documents(
                    knowledge_item_id, stored_name, original_name, media_type, size_bytes
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                (item_id, stored_name, original_name, media_type, len(data)),
            )
            chunk_count = _replace_chunks(connection, item_id, title, extracted)
    except Exception:
        target.unlink(missing_ok=True)
        raise

    return {
        "created": True,
        "duplicate": False,
        "id": item_id,
        "title": title,
        "original_name": original_name,
        "size_bytes": len(data),
        "chunk_count": chunk_count,
    }


def _item_rows(connection, item_ids: list[int] | None = None, limit: int = 250):
    where = ""
    params: list[Any] = []
    if item_ids is not None:
        if not item_ids:
            return []
        placeholders = ",".join("?" for _ in item_ids)
        where = f"WHERE ki.id IN ({placeholders})"
        params.extend(item_ids)
    params.append(limit)
    return connection.execute(
        f"""
        SELECT
            ki.id,
            ki.title,
            ki.kind,
            CASE WHEN ki.kind='watched_document' THEN NULL ELSE ki.source_path END AS source_path,
            ki.content,
            ki.created_at,
            ki.updated_at,
            kd.original_name,
            kd.media_type,
            kd.size_bytes,
            kd.imported_at,
            (
                SELECT COUNT(*) FROM knowledge_chunks kc
                WHERE kc.knowledge_item_id=ki.id
            ) AS chunk_count
        FROM knowledge_items ki
        LEFT JOIN knowledge_documents kd ON kd.knowledge_item_id=ki.id
        {where}
        ORDER BY ki.id DESC
        LIMIT ?
        """,
        tuple(params),
    ).fetchall()


def _fts_expression(query: str) -> str:
    tokens = [token for token in re.findall(r"\w+", query, flags=re.UNICODE) if token]
    tokens = tokens[:12]
    return " AND ".join(f'"{token.replace(chr(34), "")}"' for token in tokens)


def list_knowledge(query: str = "", limit: int = 250) -> list[dict[str, Any]]:
    query = query.strip()
    with db() as connection:
        if not query:
            return [dict(row) | {"snippet": None} for row in _item_rows(connection, limit=limit)]

        expression = _fts_expression(query)
        if not expression:
            term = f"%{query}%"
            rows = connection.execute(
                """
                SELECT id FROM knowledge_items
                WHERE title LIKE ? OR content LIKE ?
                ORDER BY id DESC LIMIT ?
                """,
                (term, term, limit),
            ).fetchall()
            ids = [int(row["id"]) for row in rows]
            items = [dict(row) for row in _item_rows(connection, ids, limit=limit)]
            by_id = {item["id"]: item for item in items}
            return [by_id[item_id] | {"snippet": None} for item_id in ids if item_id in by_id]

        matches = connection.execute(
            """
            SELECT
                CAST(knowledge_item_id AS INTEGER) AS knowledge_item_id,
                snippet(knowledge_chunks_fts, 0, '', '', ' … ', 32) AS snippet,
                bm25(knowledge_chunks_fts) AS score
            FROM knowledge_chunks_fts
            WHERE knowledge_chunks_fts MATCH ?
            ORDER BY score
            LIMIT ?
            """,
            (expression, max(limit * 4, 50)),
        ).fetchall()

        ordered_ids: list[int] = []
        snippets: dict[int, str] = {}
        for match in matches:
            item_id = int(match["knowledge_item_id"])
            if item_id in snippets:
                continue
            ordered_ids.append(item_id)
            snippets[item_id] = match["snippet"] or ""
            if len(ordered_ids) >= limit:
                break

        items = [dict(row) for row in _item_rows(connection, ordered_ids, limit=limit)]
        by_id = {item["id"]: item for item in items}
        return [
            by_id[item_id] | {"snippet": snippets.get(item_id)}
            for item_id in ordered_ids
            if item_id in by_id
        ]


def _attachment_paths(metadata_json: str | None) -> list[Path]:
    try:
        metadata = json.loads(str(metadata_json or "{}"))
    except (TypeError, ValueError, json.JSONDecodeError):
        return []
    if not isinstance(metadata, dict) or not isinstance(metadata.get("assets"), list):
        return []

    root = settings.knowledge_files_dir.resolve()
    attachments_root = (root / "attachments").resolve()
    paths: list[Path] = []
    for asset in metadata["assets"][:500]:
        if not isinstance(asset, dict):
            continue
        stored_name = str(asset.get("stored_name") or "").strip().replace("\\", "/")
        if not stored_name.startswith("attachments/"):
            continue
        relative = Path(stored_name)
        if relative.is_absolute() or ".." in relative.parts:
            continue
        target = (root / relative).resolve()
        if target.parent != attachments_root:
            continue
        paths.append(target)
    return paths


def delete_knowledge_item(item_id: int) -> bool:
    stored_name: str | None = None
    attachment_paths: list[Path] = []
    with db() as connection:
        item = connection.execute(
            "SELECT metadata_json FROM knowledge_items WHERE id=? LIMIT 1",
            (item_id,),
        ).fetchone()
        if item is None:
            return False
        attachment_paths = _attachment_paths(item["metadata_json"])

        document = connection.execute(
            "SELECT stored_name FROM knowledge_documents WHERE knowledge_item_id=?",
            (item_id,),
        ).fetchone()
        if document is not None:
            stored_name = document["stored_name"]

        cursor = connection.execute("DELETE FROM knowledge_items WHERE id=?", (item_id,))
        if cursor.rowcount == 0:
            return False

    if stored_name:
        try:
            (settings.knowledge_files_dir / stored_name).unlink(missing_ok=True)
        except OSError:
            pass
    for target in attachment_paths:
        try:
            target.unlink(missing_ok=True)
        except OSError:
            pass
    return True


# HomeServer v2.4 Section 3 — federated Knowledge continuity.
_SAFE_FEDERATED_KIND = re.compile(r"^[a-z0-9_.:-]{1,40}$")
_MUTATION_ID = re.compile(r"^[A-Za-z0-9._:-]{8,128}$")
_CANONICAL_ID = re.compile(r"^fd24_[0-9a-f]{40}$")
_REVISION = re.compile(r"^[0-9a-f]{64}$")


class FederatedKnowledgeError(RuntimeError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


def _knowledge_authority_key(item_id: int) -> str:
    value = int(item_id)
    if value < 1:
        raise FederatedKnowledgeError("Knowledge identity is invalid.", 500)
    return f"knowledge_item:{value}"


def _knowledge_state(item_id: int) -> dict[str, Any] | None:
    with db() as connection:
        row = connection.execute(
            """
            SELECT ki.id,ki.title,ki.kind,ki.source_path,COALESCE(ki.content,'') AS content,
                   COALESCE(ki.content_hash,'') AS content_hash,ki.created_at,ki.updated_at,
                   COALESCE(kc.collection_key,'general') AS collection_key,
                   CASE WHEN ksf.knowledge_item_id IS NULL THEN 'local_item' ELSE 'watched_folder' END AS source_type
            FROM knowledge_items ki
            LEFT JOIN knowledge_collection_items kci ON kci.knowledge_item_id=ki.id
            LEFT JOIN knowledge_collections kc ON kc.id=kci.collection_id
            LEFT JOIN knowledge_source_files ksf ON ksf.knowledge_item_id=ki.id
            WHERE ki.id=? LIMIT 1
            """,
            (int(item_id),),
        ).fetchone()
    return dict(row) if row else None


def _knowledge_revision(record: dict[str, Any]) -> str:
    content_hash = str(record.get("content_hash") or "").strip().lower()
    if not re.fullmatch(r"[0-9a-f]{64}", content_hash):
        content_hash = hashlib.sha256(
            _normalize_text(str(record.get("content") or "")).encode("utf-8")
        ).hexdigest()
    payload = {
        "title": str(record.get("title") or "").strip()[:240],
        "kind": str(record.get("kind") or "").strip().lower()[:40],
        "content_hash": content_hash,
        "collection_key": str(record.get("collection_key") or "general").strip()[:64],
        "source_type": str(record.get("source_type") or "local_item").strip()[:40],
        "updated_at": str(record.get("updated_at") or "")[:80],
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def federated_knowledge_item(record: dict[str, Any], *, snippet: str | None = None) -> dict[str, Any]:
    item_id = int(record.get("id") or 0)
    key = _knowledge_authority_key(item_id)
    title = str(record.get("title") or "Knowledge item").strip()[:240]
    content = _normalize_text(str(record.get("content") or ""))
    collection_key = str(record.get("collection_key") or "general").strip()[:64] or "general"
    source_type = str(record.get("source_type") or "local_item").strip()[:40]
    envelope = federated_data.envelope(
        "homeserver",
        "knowledge",
        key,
        title=title,
        content=content,
        updated_at=str(record.get("updated_at") or ""),
    )
    revision = _knowledge_revision(record)
    observed = dict(envelope)
    observed["record_revision"] = revision
    federated_data.observe(observed, observed_source="homeserver")
    return {
        "id": item_id,
        "title": title,
        "kind": str(record.get("kind") or "note").strip().lower()[:40],
        "snippet": None if snippet is None else str(snippet)[:6000],
        "collection_key": collection_key,
        "source_type": source_type,
        "updated_at": record.get("updated_at"),
        "authority_source": "homeserver",
        "authority_key": key,
        "canonical_id": envelope["canonical_id"],
        "record_revision": revision,
        "federation_version": federated_data.FEDERATED_DATA_VERSION,
        "mirror_only": False,
        "read_only": source_type != "local_item" or bool(str(record.get("source_path") or "").strip()),
        "mutation_route": (
            "homeserver_governed"
            if source_type == "local_item" and not str(record.get("source_path") or "").strip()
            else "source_managed_read_only"
        ),
        "allowed_mutations": (
            ["update", "delete"]
            if source_type == "local_item" and not str(record.get("source_path") or "").strip()
            else []
        ),
    }


def get_federated_knowledge_item(item_id: int) -> dict[str, Any] | None:
    row = _knowledge_state(item_id)
    return federated_knowledge_item(row) if row else None


def _item_id_from_canonical(canonical_id_value: str) -> int:
    canonical = str(canonical_id_value or "").strip()
    if not _CANONICAL_ID.fullmatch(canonical):
        raise FederatedKnowledgeError("HomeServer Knowledge canonical identity is invalid.", 422)
    key = federated_data.resolve_authority_key(
        canonical,
        authority_source="homeserver",
        dataset="knowledge",
        observed_source="homeserver",
    )
    if not key or not key.startswith("knowledge_item:"):
        raise FederatedKnowledgeError("HomeServer Knowledge item not found.", 404)
    try:
        item_id = int(key.split(":", 1)[1])
    except (TypeError, ValueError) as exc:
        raise FederatedKnowledgeError("HomeServer Knowledge identity is invalid.", 404) from exc
    expected = federated_data.canonical_id("homeserver", "knowledge", _knowledge_authority_key(item_id))
    if expected != canonical:
        raise FederatedKnowledgeError("HomeServer Knowledge identity does not match its authority.", 409)
    return item_id


def get_federated_knowledge_by_canonical(canonical_id_value: str) -> dict[str, Any] | None:
    try:
        item_id = _item_id_from_canonical(canonical_id_value)
    except FederatedKnowledgeError as exc:
        if exc.status_code == 404:
            return None
        raise
    row = _knowledge_state(item_id)
    return federated_knowledge_item(row) if row else None


def _mutation_id(value: Any) -> str:
    mutation_id = str(value or "").strip()
    if not _MUTATION_ID.fullmatch(mutation_id):
        raise FederatedKnowledgeError("Knowledge mutation_id must be 8 to 128 safe characters.")
    return mutation_id


def _expected_revision(value: Any) -> str:
    revision = str(value or "").strip().lower()
    if not _REVISION.fullmatch(revision):
        raise FederatedKnowledgeError("Knowledge expected_revision must be a SHA-256 value.")
    return revision


def _clean_kind(value: Any) -> str:
    kind = str(value or "note").strip().lower()
    if not _SAFE_FEDERATED_KIND.fullmatch(kind):
        raise FederatedKnowledgeError("Knowledge kind is invalid.")
    if kind in {"document", "watched_document"}:
        raise FederatedKnowledgeError("File-backed Knowledge kinds are source-managed and cannot be created through federated text mutations.", 409)
    return kind


def _clean_collection_key(value: Any) -> str:
    from . import knowledge_collections
    try:
        return knowledge_collections.normalize_collection_key(str(value or "general"))
    except knowledge_collections.KnowledgeCollectionError as exc:
        raise FederatedKnowledgeError(str(exc), exc.status_code) from exc


def normalize_knowledge_create_arguments(payload: dict[str, Any] | None) -> dict[str, Any]:
    raw = dict(payload or {})
    unknown = set(raw) - {"mutation_id", "collection_key", "title", "content", "kind"}
    if unknown:
        raise FederatedKnowledgeError(f"Unsupported knowledge.create argument: {sorted(unknown)[0]}")
    title = str(raw.get("title") or "").strip()
    content = _normalize_text(str(raw.get("content") or ""))
    if not title:
        raise FederatedKnowledgeError("knowledge.create requires title.")
    if len(title) > 240:
        raise FederatedKnowledgeError("Knowledge title exceeds 240 characters.")
    if not content:
        raise FederatedKnowledgeError("knowledge.create requires content.")
    if len(content) > 250000:
        raise FederatedKnowledgeError("Knowledge content exceeds 250,000 characters.", 413)
    return {
        "mutation_id": _mutation_id(raw.get("mutation_id")),
        "collection_key": _clean_collection_key(raw.get("collection_key") or "general"),
        "title": title,
        "content": content,
        "kind": _clean_kind(raw.get("kind") or "note"),
    }


def normalize_knowledge_update_arguments(payload: dict[str, Any] | None) -> dict[str, Any]:
    raw = dict(payload or {})
    unknown = set(raw) - {
        "canonical_id", "mutation_id", "expected_revision",
        "collection_key", "title", "content", "kind",
    }
    if unknown:
        raise FederatedKnowledgeError(f"Unsupported knowledge.update argument: {sorted(unknown)[0]}")
    canonical = str(raw.get("canonical_id") or "").strip()
    if not _CANONICAL_ID.fullmatch(canonical):
        raise FederatedKnowledgeError("knowledge.update requires a valid canonical_id.")
    fields: dict[str, Any] = {}
    if "title" in raw:
        title = str(raw.get("title") or "").strip()
        if not title or len(title) > 240:
            raise FederatedKnowledgeError("Knowledge title must be 1 to 240 characters.")
        fields["title"] = title
    if "content" in raw:
        content = _normalize_text(str(raw.get("content") or ""))
        if not content:
            raise FederatedKnowledgeError("Knowledge content cannot be empty.")
        if len(content) > 250000:
            raise FederatedKnowledgeError("Knowledge content exceeds 250,000 characters.", 413)
        fields["content"] = content
    if "kind" in raw:
        fields["kind"] = _clean_kind(raw.get("kind"))
    if "collection_key" in raw:
        fields["collection_key"] = _clean_collection_key(raw.get("collection_key"))
    if not fields:
        raise FederatedKnowledgeError("knowledge.update requires at least one mutable field.")
    return {
        "canonical_id": canonical,
        "mutation_id": _mutation_id(raw.get("mutation_id")),
        "expected_revision": _expected_revision(raw.get("expected_revision")),
        **fields,
    }


def normalize_knowledge_delete_arguments(payload: dict[str, Any] | None) -> dict[str, Any]:
    raw = dict(payload or {})
    unknown = set(raw) - {"canonical_id", "mutation_id", "expected_revision"}
    if unknown:
        raise FederatedKnowledgeError(f"Unsupported knowledge.delete argument: {sorted(unknown)[0]}")
    canonical = str(raw.get("canonical_id") or "").strip()
    if not _CANONICAL_ID.fullmatch(canonical):
        raise FederatedKnowledgeError("knowledge.delete requires a valid canonical_id.")
    return {
        "canonical_id": canonical,
        "mutation_id": _mutation_id(raw.get("mutation_id")),
        "expected_revision": _expected_revision(raw.get("expected_revision")),
    }


def safe_knowledge_mutation_meta(action: str, payload: dict[str, Any] | None) -> dict[str, Any]:
    raw = dict(payload or {})
    return {
        "action": str(action or "")[:32],
        "canonical_id_present": bool(str(raw.get("canonical_id") or "")),
        "canonical_id_prefix": str(raw.get("canonical_id") or "")[:5],
        "mutation_id_present": bool(str(raw.get("mutation_id") or "")),
        "expected_revision_present": bool(str(raw.get("expected_revision") or "")),
        "title_length": len(str(raw.get("title") or "")),
        "content_length": len(str(raw.get("content") or "")),
        "kind": str(raw.get("kind") or "")[:40] or None,
        "collection_key": str(raw.get("collection_key") or "")[:64] or None,
        "argument_count": len(raw),
    }


def _mutation_hash(action_key: str, payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        {"action": action_key, "payload": payload},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _mutation_replay(source_app_key: str, mutation_id: str, action_key: str, request_hash: str) -> dict[str, Any] | None:
    with db() as connection:
        row = connection.execute(
            """
            SELECT action_key,request_hash,canonical_id,result_json
            FROM federated_knowledge_mutations
            WHERE source_app_key=? AND mutation_id=? LIMIT 1
            """,
            (source_app_key, mutation_id),
        ).fetchone()
    if row is None:
        return None
    if str(row["action_key"]) != action_key or str(row["request_hash"]) != request_hash:
        raise FederatedKnowledgeError("Knowledge mutation_id was already used with different arguments.", 409)
    try:
        result = json.loads(row["result_json"] or "{}")
    except (TypeError, ValueError, json.JSONDecodeError):
        result = {}
    if not isinstance(result, dict):
        result = {}
    result["idempotent_replay"] = True
    return result


def _record_mutation(
    source_app_key: str,
    mutation_id: str,
    action_key: str,
    request_hash: str,
    canonical_id_value: str | None,
    result: dict[str, Any],
) -> None:
    with db() as connection:
        connection.execute(
            """
            INSERT INTO federated_knowledge_mutations(
                source_app_key,mutation_id,action_key,request_hash,canonical_id,result_json
            ) VALUES (?,?,?,?,?,?)
            """,
            (
                source_app_key,
                mutation_id,
                action_key,
                request_hash,
                canonical_id_value,
                json.dumps(result, ensure_ascii=False, separators=(",", ":")),
            ),
        )


def _assert_mutable_direct_item(item: dict[str, Any]) -> None:
    if bool(item.get("read_only")) or str(item.get("mutation_route") or "") != "homeserver_governed":
        raise FederatedKnowledgeError(
            "This Knowledge item is source-managed and cannot be mutated through federated continuity.",
            409,
        )


def _assert_expected_revision(item: dict[str, Any], expected_revision: str) -> None:
    current = str(item.get("record_revision") or "").strip().lower()
    if not current or current != expected_revision:
        raise FederatedKnowledgeError(
            "HomeServer Knowledge changed after this edit was prepared. Refresh and try again.",
            409,
        )


def create_federated_knowledge(
    payload: dict[str, Any],
    *,
    source_app_key: str = "owner",
) -> dict[str, Any]:
    from . import knowledge_collections
    normalized = normalize_knowledge_create_arguments(payload)
    mutation_id = str(normalized.pop("mutation_id"))
    request_hash = _mutation_hash("knowledge.create", normalized)
    replay = _mutation_replay(source_app_key, mutation_id, "knowledge.create", request_hash)
    if replay is not None:
        item = replay.get("knowledge")
        if isinstance(item, dict):
            return item
        raise FederatedKnowledgeError("Stored Knowledge mutation result is unavailable.", 500)

    collection_key = str(normalized.pop("collection_key"))
    created = create_knowledge_item(
        str(normalized["title"]),
        str(normalized["kind"]),
        str(normalized["content"]),
        None,
    )
    item_id = int(created["id"])
    try:
        knowledge_collections.assign_item(item_id, collection_key)
    except knowledge_collections.KnowledgeCollectionError as exc:
        delete_knowledge_item(item_id)
        raise FederatedKnowledgeError(str(exc), exc.status_code) from exc
    item = get_federated_knowledge_item(item_id)
    if item is None:
        raise FederatedKnowledgeError("Created HomeServer Knowledge item is unavailable.", 500)
    _record_mutation(
        source_app_key, mutation_id, "knowledge.create", request_hash,
        str(item["canonical_id"]), {"knowledge": item},
    )
    return item


def update_federated_knowledge(
    canonical_id_value: str,
    payload: dict[str, Any],
    *,
    source_app_key: str = "owner",
) -> dict[str, Any]:
    from . import knowledge_collections
    normalized = normalize_knowledge_update_arguments(
        {"canonical_id": canonical_id_value, **dict(payload or {})}
    )
    canonical = str(normalized.pop("canonical_id"))
    mutation_id = str(normalized.pop("mutation_id"))
    expected_revision = str(normalized.pop("expected_revision"))
    request_hash = _mutation_hash(
        "knowledge.update",
        {"canonical_id": canonical, "expected_revision": expected_revision, **normalized},
    )
    replay = _mutation_replay(source_app_key, mutation_id, "knowledge.update", request_hash)
    if replay is not None:
        item = replay.get("knowledge")
        if isinstance(item, dict):
            return item
        raise FederatedKnowledgeError("Stored Knowledge mutation result is unavailable.", 500)

    item_id = _item_id_from_canonical(canonical)
    current = get_federated_knowledge_item(item_id)
    if current is None:
        raise FederatedKnowledgeError("HomeServer Knowledge item not found.", 404)
    _assert_mutable_direct_item(current)
    _assert_expected_revision(current, expected_revision)
    state = _knowledge_state(item_id)
    if state is None:
        raise FederatedKnowledgeError("HomeServer Knowledge item not found.", 404)

    title = str(normalized.get("title", state["title"]))
    kind = str(normalized.get("kind", state["kind"]))
    content = _normalize_text(str(normalized.get("content", state["content"])))
    digest = hashlib.sha256(content.encode("utf-8")).hexdigest() if content else None
    with db() as connection:
        connection.execute(
            """
            UPDATE knowledge_items
            SET title=?,kind=?,content=?,content_hash=?,updated_at=CURRENT_TIMESTAMP
            WHERE id=?
            """,
            (title, kind, content, digest, item_id),
        )
        _replace_chunks(connection, item_id, title, content)
    if "collection_key" in normalized:
        try:
            knowledge_collections.assign_item(item_id, str(normalized["collection_key"]))
        except knowledge_collections.KnowledgeCollectionError as exc:
            raise FederatedKnowledgeError(str(exc), exc.status_code) from exc

    item = get_federated_knowledge_item(item_id)
    if item is None:
        raise FederatedKnowledgeError("Updated HomeServer Knowledge item is unavailable.", 500)
    _record_mutation(
        source_app_key, mutation_id, "knowledge.update", request_hash,
        canonical, {"knowledge": item},
    )
    return item


def delete_federated_knowledge(
    canonical_id_value: str,
    *,
    mutation_id: str,
    expected_revision: str,
    source_app_key: str = "owner",
) -> bool:
    canonical = str(canonical_id_value or "").strip()
    mutation_id = _mutation_id(mutation_id)
    expected_revision = _expected_revision(expected_revision)
    request_hash = _mutation_hash(
        "knowledge.delete",
        {"canonical_id": canonical, "expected_revision": expected_revision},
    )
    replay = _mutation_replay(source_app_key, mutation_id, "knowledge.delete", request_hash)
    if replay is not None:
        return bool(replay.get("deleted"))

    item_id = _item_id_from_canonical(canonical)
    current = get_federated_knowledge_item(item_id)
    if current is None:
        raise FederatedKnowledgeError("HomeServer Knowledge item not found.", 404)
    _assert_mutable_direct_item(current)
    _assert_expected_revision(current, expected_revision)
    key = _knowledge_authority_key(item_id)
    if not delete_knowledge_item(item_id):
        return False
    federated_data.mark_tombstone(
        "homeserver",
        "knowledge",
        key,
        observed_source="homeserver",
    )
    _record_mutation(
        source_app_key, mutation_id, "knowledge.delete", request_hash,
        canonical, {"deleted": True, "canonical_id": canonical},
    )
    return True
